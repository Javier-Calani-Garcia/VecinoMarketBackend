"""Exportación genérica de reportes/documentos a CSV, XLSX, PDF, HTML y
eMail — usado por los dashboards (CU18/CU19) y por las facturas (CU26).

Formato común de entrada para no repetir la generación de cada tipo de
archivo por cada consumidor: una lista de "secciones", cada una con un
título, encabezados de columna y filas de datos ya formateados a texto.
"""
import csv
import html
import io
from datetime import datetime
from xml.sax.saxutils import escape as _escape_xml

from django.conf import settings
from django.core.mail import EmailMessage
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

COLOR_MARCA = '#D97706'
FORMATOS_VALIDOS = ('csv', 'xlsx', 'pdf', 'html', 'email')

_CONTENT_TYPES = {
    'csv': 'text/csv; charset=utf-8',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'pdf': 'application/pdf',
    'html': 'text/html; charset=utf-8',
}


def _nombre_archivo(base, extension):
    fecha = datetime.now().strftime('%Y%m%d_%H%M')
    return f'{base}_{fecha}.{extension}'


def _generar_bytes(formato, titulo, subtitulo, secciones):
    if formato == 'csv':
        return _csv_bytes(titulo, subtitulo, secciones)
    if formato == 'xlsx':
        return _xlsx_bytes(titulo, subtitulo, secciones)
    if formato == 'pdf':
        return _pdf_bytes(titulo, subtitulo, secciones)
    if formato == 'html':
        return _html_bytes(titulo, subtitulo, secciones)
    raise ValueError(f'Formato de exportación no soportado: {formato}')


def exportar_reporte(formato, base_nombre, titulo, subtitulo, secciones):
    """secciones: [{'titulo': str, 'headers': [...], 'filas': [[...], ...]}]"""
    contenido = _generar_bytes(formato, titulo, subtitulo, secciones)
    response = HttpResponse(contenido, content_type=_CONTENT_TYPES[formato])
    response['Content-Disposition'] = f'attachment; filename="{_nombre_archivo(base_nombre, formato)}"'
    return response


def responder_exportacion(request, base_nombre, titulo, subtitulo, secciones):
    """Punto único usado por todas las vistas de exportación (dashboards,
    facturas): valida el ?formato=, y si es 'email' lo manda al correo del
    usuario en vez de descargarlo (correo_recuperacion si existe — mismo
    criterio que las cuentas de empresa, ver apps/usuarios/services.py)."""
    from rest_framework import status as _status
    from rest_framework.response import Response as _Response

    formato = request.query_params.get('formato', 'pdf').lower()
    if formato not in FORMATOS_VALIDOS:
        return _Response(
            {'detail': 'Formato inválido. Usa csv, xlsx, pdf, html o email.'}, status=_status.HTTP_400_BAD_REQUEST
        )

    if formato == 'email':
        destinatario = getattr(request.user, 'correo_recuperacion', '') or request.user.email
        enviar_reporte_por_email(destinatario, base_nombre, titulo, subtitulo, secciones)
        return _Response({'detail': f'Reporte enviado a {destinatario}.'})

    return exportar_reporte(formato, base_nombre, titulo, subtitulo, secciones)


def enviar_reporte_por_email(destinatario, base_nombre, titulo, subtitulo, secciones, formato_adjunto='pdf'):
    """Manda el reporte como adjunto al correo del usuario que lo pidió,
    en vez de descargarlo — la opción "eMail" de exportación."""
    contenido = _generar_bytes(formato_adjunto, titulo, subtitulo, secciones)
    nombre_archivo = _nombre_archivo(base_nombre, formato_adjunto)

    correo = EmailMessage(
        f'{titulo} — VecinoMarket',
        f'Adjunto encontrarás el reporte solicitado.\n\n{subtitulo}',
        settings.DEFAULT_FROM_EMAIL,
        [destinatario],
    )
    correo.attach(nombre_archivo, contenido, _CONTENT_TYPES[formato_adjunto])
    correo.send(fail_silently=False)


def _csv_bytes(titulo, subtitulo, secciones):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([titulo])
    writer.writerow([subtitulo])
    for seccion in secciones:
        writer.writerow([])
        writer.writerow([seccion['titulo']])
        writer.writerow(seccion['headers'])
        for fila in seccion['filas']:
            writer.writerow(fila)

    contenido = '﻿' + buffer.getvalue()  # BOM para que Excel abra bien los acentos
    return contenido.encode('utf-8')


def _e(valor):
    """Escapa cualquier valor para insertarlo en HTML -- los reportes dinámicos
    incluyen texto libre real (nombres de producto, de cliente, etc.) que puede
    traer '&', '<' o '>' y romper el marcado si no se escapa."""
    return html.escape(str(valor), quote=False)


def _pe(valor):
    """Igual que _e() pero para el mini-XML que interpreta Paragraph de
    reportlab (PDF) -- exige str(): muchas celdas de reportes vienen como
    int/float/None (conteos, montos) y xml.sax.saxutils.escape solo acepta
    strings."""
    return _escape_xml(str(valor))


def _html_bytes(titulo, subtitulo, secciones):
    generado = timezone.now().strftime('%d/%m/%Y %H:%M')
    partes = [
        '<!doctype html><html lang="es"><head><meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width, initial-scale=1">',
        f'<title>{_e(titulo)}</title>',
        '<style>',
        ':root{color-scheme:light;}',
        'body{font-family:-apple-system,Segoe UI,Roboto,Arial,Helvetica,sans-serif;'
        'color:#1f2937;background:#f3f4f6;margin:0;padding:24px 16px;}',
        '.hoja{max-width:1000px;margin:0 auto;background:#fff;border:1px solid #e5e7eb;'
        'border-radius:12px;padding:28px 32px;box-shadow:0 1px 3px rgba(0,0,0,.06);}',
        f'.marca{{color:{COLOR_MARCA};font-size:12px;font-weight:700;letter-spacing:.06em;'
        'text-transform:uppercase;margin:0 0 6px;}}',
        'h1{margin:0 0 4px;font-size:22px;color:#111827;}',
        '.subtitulo{color:#6b7280;margin:0 0 22px;font-size:13px;}',
        '.seccion{margin-top:26px;}',
        'h2{font-size:15px;margin:0 0 10px;color:#111827;border-left:4px solid ' + COLOR_MARCA + ';padding-left:8px;}',
        '.tabla-wrap{overflow-x:auto;border:1px solid #e5e7eb;border-radius:8px;}',
        'table{border-collapse:collapse;width:100%;min-width:480px;font-size:13px;}',
        f'th{{background:{COLOR_MARCA};color:#fff;text-align:left;padding:8px 12px;'
        'white-space:nowrap;position:sticky;top:0;}}',
        'td{padding:7px 12px;border-bottom:1px solid #f0f0f0;color:#374151;}',
        'tr:last-child td{border-bottom:none;}',
        'tr:nth-child(even) td{background:#fafafa;}',
        '.sin-datos{color:#9ca3af;font-style:italic;}',
        '.pie{margin-top:28px;padding-top:14px;border-top:1px solid #e5e7eb;'
        'color:#9ca3af;font-size:11px;}',
        '@media print{body{background:#fff;padding:0;}.hoja{box-shadow:none;border:none;}}',
        '</style></head><body>',
        '<div class="hoja">',
        '<p class="marca">VecinoMarket</p>',
        f'<h1>{_e(titulo)}</h1>',
        f'<p class="subtitulo">{_e(subtitulo)}</p>',
    ]
    for seccion in secciones:
        partes.append('<div class="seccion">')
        partes.append(f'<h2>{_e(seccion["titulo"])}</h2>')
        filas = seccion['filas']
        if not filas:
            partes.append('<p class="sin-datos">Sin datos para este rango/filtro.</p>')
        else:
            partes.append('<div class="tabla-wrap"><table><thead><tr>')
            partes.extend(f'<th>{_e(h)}</th>' for h in seccion['headers'])
            partes.append('</tr></thead><tbody>')
            for fila in filas:
                partes.append('<tr>' + ''.join(f'<td>{_e(v)}</td>' for v in fila) + '</tr>')
            partes.append('</tbody></table></div>')
        partes.append('</div>')
    partes.append(f'<p class="pie">Generado el {generado} · VecinoMarket</p>')
    partes.append('</div></body></html>')
    return ''.join(partes).encode('utf-8')


def _xlsx_bytes(titulo, subtitulo, secciones):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Reporte'

    fila_actual = 1
    ws.cell(row=fila_actual, column=1, value=titulo).font = Font(size=14, bold=True, color='D97706')
    fila_actual += 1
    ws.cell(row=fila_actual, column=1, value=subtitulo).font = Font(italic=True, color='666666')
    fila_actual += 2

    for seccion in secciones:
        ws.cell(row=fila_actual, column=1, value=seccion['titulo']).font = Font(bold=True, size=12)
        fila_actual += 1

        for col, encabezado in enumerate(seccion['headers'], start=1):
            celda = ws.cell(row=fila_actual, column=col, value=encabezado)
            celda.font = Font(bold=True, color='FFFFFF')
            celda.fill = PatternFill('solid', fgColor=COLOR_MARCA.lstrip('#'))
            celda.alignment = Alignment(horizontal='left')
        fila_actual += 1

        filas = seccion['filas'] or [['Sin datos']]
        for fila in filas:
            for col, valor in enumerate(fila, start=1):
                ws.cell(row=fila_actual, column=col, value=valor)
            fila_actual += 1
        fila_actual += 1  # espacio entre secciones

    for col in range(1, 6):
        ws.column_dimensions[get_column_letter(col)].width = 26

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(titulo, subtitulo, secciones):
    buffer = io.BytesIO()
    # Landscape: los reportes dinámicos pueden traer bastantes columnas (el
    # usuario elige cuáles incluir) y en vertical se salían de la hoja.
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(letter),
        topMargin=1.3 * cm, bottomMargin=1.3 * cm, leftMargin=1.3 * cm, rightMargin=1.3 * cm,
    )
    estilos = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle('TituloReporte', parent=estilos['Heading1'], textColor=colors.HexColor(COLOR_MARCA))
    estilo_subtitulo = ParagraphStyle('Subtitulo', parent=estilos['Normal'], textColor=colors.HexColor('#666666'), spaceAfter=14)
    estilo_seccion = ParagraphStyle('Seccion', parent=estilos['Heading2'], spaceBefore=16, spaceAfter=6, fontSize=13)
    estilo_pie = ParagraphStyle('Pie', parent=estilos['Normal'], textColor=colors.HexColor('#9ca3af'), fontSize=8, spaceBefore=18)

    elementos = [Paragraph(_pe(titulo), estilo_titulo), Paragraph(_pe(subtitulo), estilo_subtitulo)]

    for seccion in secciones:
        elementos.append(Paragraph(_pe(seccion['titulo']), estilo_seccion))

        filas = seccion['filas']
        if not filas:
            elementos.append(Paragraph('Sin datos para este rango/filtro.', ParagraphStyle(
                'SinDatos', parent=estilos['Normal'], textColor=colors.HexColor('#9ca3af'), fontName='Helvetica-Oblique',
            )))
            elementos.append(Spacer(1, 4))
            continue

        num_columnas = len(seccion['headers'])
        # Texto más chico cuantas más columnas haya, para que quepan sin
        # amontonarse ni desbordar la página.
        tamano_fuente = 8.5 if num_columnas <= 5 else (7.5 if num_columnas <= 8 else 6.5)
        estilo_celda = ParagraphStyle('Celda', parent=estilos['Normal'], fontSize=tamano_fuente, leading=tamano_fuente + 2.5)
        estilo_encabezado = ParagraphStyle('Encabezado', parent=estilo_celda, textColor=colors.white, fontName='Helvetica-Bold')

        encabezados = [Paragraph(_pe(h), estilo_encabezado) for h in seccion['headers']]
        cuerpo = [[Paragraph(_pe(v), estilo_celda) for v in fila] for fila in filas]
        datos_tabla = [encabezados] + cuerpo

        # Ancho de columna proporcional al contenido real (encabezado y una
        # muestra de filas) -- si solo mirara el encabezado, una columna con
        # encabezado corto pero datos largos (ej. "Fecha") quedaba angosta y
        # el texto se cortaba feo a la mitad.
        muestra = filas[:200]
        anchos_relativos = [
            min(max(len(str(seccion['headers'][i])), *(len(str(fila[i])) for fila in muestra), 6), 40)
            for i in range(num_columnas)
        ]
        total_relativo = sum(anchos_relativos)
        ancho_disponible = doc.width
        anchos = [ancho_disponible * (r / total_relativo) for r in anchos_relativos]

        tabla = Table(datos_tabla, colWidths=anchos, hAlign='LEFT', repeatRows=1)
        tabla.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_MARCA)),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#DDDDDD')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FAFAFA')]),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ]))
        elementos.append(tabla)
        elementos.append(Spacer(1, 10))

    elementos.append(Paragraph(f'Generado el {timezone.now().strftime("%d/%m/%Y %H:%M")} · VecinoMarket', estilo_pie))

    doc.build(elementos)
    return buffer.getvalue()
